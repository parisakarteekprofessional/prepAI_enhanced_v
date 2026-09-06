import os
import re
from typing import List, Dict, Any

def extract_pdf_pages(file_path: str) -> List[Dict[str, Any]]:
    import pymupdf  # PyMuPDF
    doc = pymupdf.open(file_path)
    pages = []
    total_printable_chars = 0

    for idx, page in enumerate(doc):
        text = page.get_text("text") or ""
        clean_text = text.strip()
        # Count alphanumeric characters to detect scanned / empty PDFs
        alphanumeric_count = len(re.findall(r'[a-zA-Z0-9]', clean_text))
        total_printable_chars += alphanumeric_count

        pages.append({
            "pageNumber": idx + 1,
            "text": clean_text
        })

    doc.close()

    if len(pages) == 0:
        raise ValueError("The uploaded PDF document is completely empty.")

    if total_printable_chars < 20:
        raise ValueError(
            "Scanned or image-only PDF detected without selectable text. "
            "Please upload a searchable document containing selectable text."
        )

    return pages


def extract_docx_pages(file_path: str) -> List[Dict[str, Any]]:
    import docx
    doc = docx.Document(file_path)
    paragraphs = []
    current_page_text = []
    pages = []
    char_count = 0

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        current_page_text.append(txt)
        char_count += len(txt)
        # Approximate 2500 characters per page for DOCX
        if char_count >= 2500:
            pages.append({
                "pageNumber": len(pages) + 1,
                "text": "\n\n".join(current_page_text)
            })
            current_page_text = []
            char_count = 0

    # Also extract any tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                table_rows.append(row_text)
        if table_rows:
            current_page_text.append("\n".join(table_rows))

    if current_page_text:
        pages.append({
            "pageNumber": len(pages) + 1,
            "text": "\n\n".join(current_page_text)
        })

    if not pages or sum(len(p["text"]) for p in pages) < 10:
        raise ValueError("The uploaded DOCX document contains no readable text.")

    return pages


def extract_txt_or_md_pages(file_path: str) -> List[Dict[str, Any]]:
    text = ""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            text = f.read()

    clean_text = text.strip()
    if len(clean_text) < 5:
        raise ValueError("The uploaded text document is empty.")

    # Split into logical sections of approx 3000 chars
    raw_sections = clean_text.split("\n\n")
    pages = []
    cur_chunk = []
    cur_len = 0

    for sec in raw_sections:
        s = sec.strip()
        if not s:
            continue
        cur_chunk.append(s)
        cur_len += len(s)
        if cur_len >= 3000:
            pages.append({
                "pageNumber": len(pages) + 1,
                "text": "\n\n".join(cur_chunk)
            })
            cur_chunk = []
            cur_len = 0

    if cur_chunk:
        pages.append({
            "pageNumber": len(pages) + 1,
            "text": "\n\n".join(cur_chunk)
        })

    return pages


def extract_document_content(file_path: str, file_type: str) -> List[Dict[str, Any]]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found on disk: {file_path}")

    norm_type = file_type.lower().replace(".", "")

    if norm_type == "pdf":
        return extract_pdf_pages(file_path)
    elif norm_type == "docx":
        return extract_docx_pages(file_path)
    elif norm_type in ["txt", "md"]:
        return extract_txt_or_md_pages(file_path)
    else:
        raise ValueError(f"Unsupported document format: .{norm_type}. Supported types: PDF, DOCX, TXT, MD.")
